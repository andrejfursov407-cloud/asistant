import os
import sys
import json
import threading
import html
from html.parser import HTMLParser
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import customtkinter as ctk
import requests
from openai import OpenAI
import PyPDF2
import docx

APP_NAME = "Помощник AI"
CONFIG_FILE = Path.home() / ".pomoshchnik_ai_config.json"
MAX_CONTEXT_CHARS = 60000
ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")


class SearchParser(HTMLParser):
    """Extract readable text and links from DuckDuckGo HTML results."""
    def __init__(self):
        super().__init__()
        self.results = []
        self._in_result = False
        self._in_snippet = False
        self._current = {"title": "", "url": "", "snippet": ""}
        self._capture = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        classes = attrs.get("class", "")
        if tag == "a" and "result__a" in classes:
            self._current = {"title": "", "url": attrs.get("href", ""), "snippet": ""}
            self._capture = "title"
        elif tag in ("a", "div") and ("result__snippet" in classes):
            self._capture = "snippet"

    def handle_endtag(self, tag):
        if tag == "a" and self._capture == "title":
            if self._current.get("title"):
                self.results.append(self._current.copy())
            self._capture = None
        elif tag in ("a", "div") and self._capture == "snippet":
            self._capture = None

    def handle_data(self, data):
        if self._capture in ("title", "snippet"):
            self._current[self._capture] += data.strip() + " "


class PomoshchnikAI(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title(APP_NAME)
        self.geometry("1120x760")
        self.minsize(900, 620)
        self.document_text = ""
        self.document_names = []
        self.load_config()

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)
        self.sidebar = ctk.CTkFrame(self, width=285, corner_radius=0)
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_columnconfigure(0, weight=1)
        self.sidebar.grid_rowconfigure(7, weight=1)

        ctk.CTkLabel(self.sidebar, text="✦ Помощник AI",
                     font=("Segoe UI", 22, "bold")).grid(
            row=0, column=0, padx=18, pady=(24, 20), sticky="w")
        ctk.CTkLabel(self.sidebar, text="Провайдер ИИ",
                     font=("Segoe UI", 12, "bold")).grid(
            row=1, column=0, padx=18, sticky="w")

        self.providers = {
            "Google Gemini (бесплатный тариф)": {
                "url": "https://generativelanguage.googleapis.com/v1beta/openai/",
                "model": "gemini-2.0-flash",
            },
            "Groq (бесплатный тариф)": {
                "url": "https://api.groq.com/openai/v1",
                "model": "llama-3.3-70b-versatile",
            },
            "OpenRouter (доступные бесплатные модели)": {
                "url": "https://openrouter.ai/api/v1",
                "model": "meta-llama/llama-3.3-70b-instruct:free",
            },
        }
        self.provider_menu = ctk.CTkOptionMenu(
            self.sidebar, values=list(self.providers), command=self.provider_changed)
        self.provider_menu.grid(row=2, column=0, padx=18, pady=(6, 14), sticky="ew")
        if self.saved_provider in self.providers:
            self.provider_menu.set(self.saved_provider)

        ctk.CTkLabel(self.sidebar, text="API-ключ выбранного сервиса",
                     font=("Segoe UI", 12, "bold")).grid(
            row=3, column=0, padx=18, sticky="w")
        self.api_entry = ctk.CTkEntry(self.sidebar, show="*", placeholder_text="Вставьте API-ключ")
        self.api_entry.grid(row=4, column=0, padx=18, pady=(6, 10), sticky="ew")
        if self.saved_api_key:
            self.api_entry.insert(0, self.saved_api_key)

        self.save_btn = ctk.CTkButton(self.sidebar, text="Сохранить настройки",
                                      command=self.save_config, fg_color="#475569")
        self.save_btn.grid(row=5, column=0, padx=18, pady=(0, 16), sticky="ew")

        ctk.CTkLabel(self.sidebar, text="Документы",
                     font=("Segoe UI", 13, "bold")).grid(
            row=6, column=0, padx=18, sticky="w")
        self.doc_status = ctk.CTkLabel(
            self.sidebar, text="Файлы не выбраны", justify="left",
            wraplength=240, anchor="w")
        self.doc_status.grid(row=7, column=0, padx=18, pady=8, sticky="new")

        self.load_docs_btn = ctk.CTkButton(
            self.sidebar, text="Добавить документы…", command=self.load_documents)
        self.load_docs_btn.grid(row=8, column=0, padx=18, pady=5, sticky="ew")
        self.save_txt_btn = ctk.CTkButton(
            self.sidebar, text="Объединить и сохранить TXT…",
            command=self.export_txt, fg_color="#0f766e")
        self.save_txt_btn.grid(row=9, column=0, padx=18, pady=5, sticky="ew")
        self.save_docx_btn = ctk.CTkButton(
            self.sidebar, text="Объединить и сохранить DOCX…",
            command=self.export_docx, fg_color="#0f766e")
        self.save_docx_btn.grid(row=10, column=0, padx=18, pady=(5, 18), sticky="ew")

        self.main = ctk.CTkFrame(self, corner_radius=0)
        self.main.grid(row=0, column=1, sticky="nsew")
        self.main.grid_columnconfigure(0, weight=1)
        self.main.grid_rowconfigure(1, weight=1)
        self.top = ctk.CTkFrame(self.main, height=56, corner_radius=0)
        self.top.grid(row=0, column=0, sticky="ew")
        ctk.CTkLabel(self.top, text="Написание текстов · Поиск · Анализ документов",
                     font=("Segoe UI", 15, "bold")).pack(side="left", padx=20, pady=15)
        self.chat = ctk.CTkTextbox(self.main, wrap="word", font=("Segoe UI", 13))
        self.chat.grid(row=1, column=0, padx=18, pady=12, sticky="nsew")
        self.chat.configure(state="disabled")

        self.actions = ctk.CTkFrame(self.main, fg_color="transparent")
        self.actions.grid(row=2, column=0, padx=18, pady=(0, 6), sticky="ew")
        self.actions.grid_columnconfigure(0, weight=1)
        self.mode = ctk.CTkOptionMenu(
            self.actions, values=[
                "Обычный ответ", "Написать текст", "Проанализировать документы",
                "Поиск в интернете и ответ"
            ], width=245)
        self.mode.set("Обычный ответ")
        self.mode.grid(row=0, column=0, sticky="w", pady=(0, 6))

        self.input = ctk.CTkTextbox(self.main, height=95, wrap="word", font=("Segoe UI", 13))
        self.input.grid(row=3, column=0, padx=18, pady=(0, 10), sticky="ew")
        self.input.bind("<Control-Return>", lambda event: self.send_message())
        self.bottom = ctk.CTkFrame(self.main, fg_color="transparent")
        self.bottom.grid(row=4, column=0, padx=18, pady=(0, 16), sticky="ew")
        self.status = ctk.CTkLabel(self.bottom, text="Готов к работе", anchor="w")
        self.status.pack(side="left", fill="x", expand=True)
        self.send_btn = ctk.CTkButton(self.bottom, text="Отправить  ➤",
                                      command=self.send_message, width=150)
        self.send_btn.pack(side="right")
        self.append("Система", "Добро пожаловать! Выберите режим, введите запрос. Можно загрузить несколько PDF, DOCX и TXT, проанализировать их и объединить в один файл.\n\nДля ИИ нужен API-ключ выбранного провайдера. Бесплатные тарифы имеют лимиты и могут меняться. Не отправляйте в сторонние сервисы конфиденциальные документы без разрешения организации.")
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def load_config(self):
        self.saved_api_key = ""
        self.saved_provider = "Google Gemini (бесплатный тариф)"
        try:
            if CONFIG_FILE.exists():
                data = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
                self.saved_api_key = data.get("api_key", "")
                self.saved_provider = data.get("provider", self.saved_provider)
        except Exception:
            pass

    def save_config(self, silent=False):
        data = {"api_key": self.api_entry.get().strip(), "provider": self.provider_menu.get()}
        try:
            CONFIG_FILE.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            if not silent:
                messagebox.showinfo("Помощник AI", f"Настройки сохранены в профиле пользователя:\n{CONFIG_FILE}")
        except Exception as exc:
            if not silent:
                messagebox.showerror("Ошибка", f"Не удалось сохранить настройки: {exc}")

    def on_close(self):
        self.save_config(silent=True)
        self.destroy()

    def provider_changed(self, _choice):
        pass

    def append(self, who, text):
        self.chat.configure(state="normal")
        self.chat.insert("end", f"\n{'─' * 44}\n{who}\n{'─' * 44}\n{text}\n")
        self.chat.configure(state="disabled")
        self.chat.see("end")

    def set_status(self, text):
        self.after(0, lambda: self.status.configure(text=text))

    def read_file(self, path):
        ext = path.suffix.lower()
        if ext == ".pdf":
            parts = []
            with path.open("rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    parts.append(page.extract_text() or "")
            return "\n".join(parts)
        if ext == ".docx":
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        if ext == ".txt":
            return path.read_text(encoding="utf-8-sig")
        raise ValueError(f"Формат {ext} не поддерживается")

    def load_documents(self):
        names = filedialog.askopenfilenames(
            title="Выберите один или несколько документов",
            filetypes=[("Документы PDF, Word, TXT", "*.pdf *.docx *.txt"),
                       ("Все файлы", "*.*")])
        if not names:
            return
        pieces, loaded, errors = [], [], []
        for name in names:
            path = Path(name)
            try:
                text = self.read_file(path).strip()
                if not text:
                    errors.append(f"{path.name}: текст не извлечён (возможно, PDF состоит из сканов)")
                    continue
                pieces.append(f"\n\n===== ДОКУМЕНТ: {path.name} =====\n{text}")
                loaded.append(path.name)
            except Exception as exc:
                errors.append(f"{path.name}: {exc}")
        if pieces:
            self.document_text = "\n".join(pieces)
            self.document_names = loaded
            shown = "\n".join("• " + n for n in loaded)
            if len(shown) > 850:
                shown = shown[:850] + "…"
            self.doc_status.configure(text=f"Загружено: {len(loaded)}\n{shown}")
            self.append("Система", f"Загружено документов: {len(loaded)}. Их текст доступен для анализа и объединения.")
        if errors:
            messagebox.showwarning("Некоторые файлы пропущены", "\n".join(errors[:10]))

    def export_txt(self):
        if not self.document_text:
            messagebox.showwarning("Нет документов", "Сначала загрузите документы.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".txt",
                    filetypes=[("Текстовый файл", "*.txt")], initialfile="объединённые_документы.txt")
        if path:
            Path(path).write_text(self.document_text, encoding="utf-8")
            messagebox.showinfo("Готово", f"Файл сохранён:\n{path}")

    def export_docx(self):
        if not self.document_text:
            messagebox.showwarning("Нет документов", "Сначала загрузите документы.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".docx",
                    filetypes=[("Документ Word", "*.docx")], initialfile="объединённые_документы.docx")
        if path:
            out = docx.Document()
            out.add_heading("Объединённые документы", 0)
            for chunk in self.document_text.split("\n\n===== ДОКУМЕНТ: "):
                if not chunk.strip():
                    continue
                lines = chunk.splitlines()
                title = lines[0].replace("=====", "").strip() if lines else "Документ"
                out.add_heading(title, level=1)
                body = "\n".join(lines[1:]) if len(lines) > 1 else chunk
                for para in body.splitlines():
                    if para.strip():
                        out.add_paragraph(para)
            out.save(path)
            messagebox.showinfo("Готово", f"Файл сохранён:\n{path}")

    def web_search(self, query):
        # Ordinary HTTPS search request; no spoofed browser headers or traffic concealment.
        response = requests.get(
            "https://html.duckduckgo.com/html/",
            params={"q": query}, timeout=20,
            headers={"User-Agent": "PomoshchnikAI/1.0 (web search)"})
        response.raise_for_status()
        parser = SearchParser()
        parser.feed(response.text)
        return parser.results[:7]

    def send_message(self):
        question = self.input.get("1.0", "end").strip()
        if not question:
            return
        api_key = self.api_entry.get().strip()
        if not api_key:
            messagebox.showwarning("Нужен API-ключ", "Введите API-ключ выбранного провайдера. Ключ выдаётся на сайте провайдера; бесплатный доступ зависит от его актуальных условий.")
            return
        self.input.delete("1.0", "end")
        self.append("Вы", question)
        self.send_btn.configure(state="disabled", text="Обработка…")
        self.set_status("Выполняется запрос…")
        provider = self.provider_menu.get()
        settings = self.providers[provider]
        mode = self.mode.get()
        threading.Thread(target=self.worker, args=(question, api_key, provider, settings, mode), daemon=True).start()

    def worker(self, question, api_key, provider, settings, mode):
        try:
            web_context = ""
            if mode == "Поиск в интернете и ответ":
                self.set_status("Поиск в интернете…")
                results = self.web_search(question)
                if results:
                    web_context = "\n\nРезультаты веб-поиска (проверяйте источники и даты):\n" + "\n".join(
                        f"- {r['title'].strip()}\n  {r['url']}\n  {r['snippet'].strip()}" for r in results)
                else:
                    web_context = "\nПоиск не вернул результатов."
            system_prompt = (
                "Ты — Помощник AI, корпоративный помощник по написанию текстов и анализу документов. "
                "Отвечай на русском языке, структурированно и по делу. Не выдумывай факты. "
                "Если опираешься на документы, указывай названия документов. "
                "Если данных недостаточно, прямо сообщи об этом."
            )
            if mode == "Написать текст":
                system_prompt += " Помоги подготовить качественный текст по цели, аудитории и формату пользователя."
            elif mode == "Проанализировать документы":
                system_prompt += " Проанализируй загруженные документы: выдели основные положения, противоречия, риски, выводы и практические рекомендации. Не приписывай документам отсутствующие сведения."
            if self.document_text:
                system_prompt += "\n\nТЕКСТ ЗАГРУЖЕННЫХ ДОКУМЕНТОВ:\n" + self.document_text[:MAX_CONTEXT_CHARS]
                if len(self.document_text) > MAX_CONTEXT_CHARS:
                    system_prompt += "\n[ВНИМАНИЕ: контекст сокращён до лимита; не весь текст передан модели.]"
            if web_context:
                system_prompt += "\n\n" + web_context
            self.set_status("Запрос к ИИ…")
            client = OpenAI(api_key=api_key, base_url=settings["url"], timeout=60, max_retries=1)
            response = client.chat.completions.create(
                model=settings["model"],
                messages=[{"role": "system", "content": system_prompt},
                          {"role": "user", "content": question}],
                temperature=0.2,
            )
            answer = response.choices[0].message.content or "(Пустой ответ)"
            self.after(0, self.append, f"ИИ · {provider}", answer)
            self.set_status("Готово")
        except Exception as exc:
            self.after(0, self.append, "Ошибка", f"{exc}\n\nПроверьте интернет-соединение, API-ключ, доступность модели и ограничения бесплатного тарифа.")
            self.set_status("Ошибка запроса")
        finally:
            self.after(0, lambda: self.send_btn.configure(state="normal", text="Отправить  ➤"))


if __name__ == "__main__":
    app = PomoshchnikAI()
    app.mainloop()
